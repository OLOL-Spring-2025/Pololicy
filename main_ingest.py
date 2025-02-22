import os
import shutil
from pathlib import Path
from pdf2image import convert_from_path
import pytesseract
import cv2
from pypdf import PdfReader
from fpdf import FPDF
import pandas as pd
from docx import Document
from docx2pdf import convert
import sys
import os
import comtypes.client

# Set up paths
INPUT_DIR = "OneDrive Feb 12 2025/Academic Affairs"
OUTPUT_DIR = "output_texts"
PDF_DIR = "converted_pdfs"

os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(PDF_DIR, exist_ok=True)

def convert_doc_to_pdf(input_path, output_path):
    wdFormatPDF = 17

    # Ensure the file exists
    if not os.path.exists(input_path):
        print(f"Error: File not found -> {input_path}")
        return

    #print(f"Converting DOC to PDF: {input_path} -> {output_path}")

    # Initialize Word application
    word = comtypes.client.CreateObject('Word.Application')
    word.Visible = False  # Run in the background

    try:
        doc = word.Documents.Open(os.path.abspath(input_path))  # Ensure absolute path
        doc.SaveAs(os.path.abspath(output_path), FileFormat=wdFormatPDF)
        doc.Close()
        #print(f"Successfully converted: {output_path}")

    except Exception as e:
        print(f"Error converting DOC to PDF: {e}")

    finally:
        word.Quit()


# Function to convert XLS/XLSX to PDF (text extraction)
def convert_xls_to_pdf(input_path, output_path):
    #print(f"Converting XLS to PDF: {input_path} -> {output_path}")
    try:
        df = pd.read_excel(input_path, sheet_name=None)
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_font("Arial", size=10)

        for sheet, data in df.items():
            pdf.cell(200, 10, txt=f"Sheet: {sheet}", ln=True)
            for row in data.astype(str).values:
                pdf.cell(200, 10, txt=" | ".join(row), ln=True)

        pdf.output(output_path)
    except Exception as e:
        print(f"Error converting XLS: {e}")

# Function to convert text files to PDF
def convert_txt_to_pdf(input_path, output_path):
    #print(f"Converting TXT to PDF: {input_path} -> {output_path}")
    try:
        with open(input_path, "r", encoding="utf-8") as file:
            content = file.readlines()
        
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_font("Arial", size=12)

        for line in content:
            pdf.cell(200, 10, txt=line.strip(), ln=True)

        pdf.output(output_path)
    except Exception as e:
        print(f"Error converting TXT: {e}")

# Function to extract text from PDFs
def extract_text_from_pdf(pdf_path):
    #print(f"Extracting text from PDF: {pdf_path}")
    try:
        reader = PdfReader(pdf_path)
        text = ""

        for page in reader.pages:
            extracted_text = page.extract_text()
            text += extracted_text if extracted_text else ""

        if not text.strip():
            print(f"No text found, using OCR... {pdf_path}")
            text = extract_text_via_ocr(pdf_path)

        return text
    except Exception as e:
        print(f"Error reading PDF: {e}")
        return ""

# Function to extract text from scanned PDFs using OCR
def extract_text_via_ocr(pdf_path):
    #print(f"Performing OCR on scanned PDF: {pdf_path}")
    text = ""
    try:
        images = convert_from_path(pdf_path)
        for img in images:
            img_cv = cv2.cvtColor(cv2.imread(img.filename), cv2.COLOR_BGR2GRAY)
            text += pytesseract.image_to_string(img_cv)
    except Exception as e:
        print(f"OCR Error: {e}")
    return text

# Function to process files in the directory
def process_files(input_dir):
    for root, _, files in os.walk(input_dir):
        for file in files:
            input_path = os.path.join(root, file)
            filename, ext = os.path.splitext(file)
            ext = ext.lower()

            #print(f"Processing file: {file} (Extension: {ext})")

            if ext == ".pdf":
                pdf_path = input_path  # Already a PDF
            else:
                pdf_path = os.path.join(PDF_DIR, f"{filename}.pdf")

                if ext in [".docx"]:
                    convert(input_path, pdf_path)
                elif ext in [".doc"]:
                    convert_doc_to_pdf(input_path, pdf_path)    
                elif ext in [".xls", ".xlsx"]:
                    convert_xls_to_pdf(input_path, pdf_path)
                elif ext in [".txt"]:
                    convert_txt_to_pdf(input_path, pdf_path)
                else:
                    print(f"Skipping unsupported file type: {file}")
                    continue

            if os.path.exists(pdf_path):
                extracted_text = extract_text_from_pdf(pdf_path)

                # Save extracted text
                text_file_path = os.path.join(OUTPUT_DIR, f"{filename}.txt")
                with open(text_file_path, "w", encoding="utf-8") as text_file:
                    text_file.write(extracted_text)

                #print(f"Text extracted and saved: {text_file_path}")
            else:
                print(f"PDF not found: {pdf_path}")

# Run the process
if __name__ == "__main__":
    if not os.path.exists(INPUT_DIR):
        print(f"Error: Input directory '{INPUT_DIR}' does not exist!")
    else:
        process_files(INPUT_DIR)
